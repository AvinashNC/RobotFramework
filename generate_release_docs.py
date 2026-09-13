#!/usr/bin/env python3
"""
generate_release_docs.py

Automates release note generation:
1. Fetches the diff (commits + file changes) between two git tags from GitHub.
2. Feeds that diff + a previous release doc (as a style reference) to Google Gemini.
3. Writes the generated release doc to a markdown file AND a matching .docx file.

Setup:
    pip install requests python-docx

    Environment variables (or pass via CLI flags):
        GITHUB_TOKEN   - GitHub personal access token (needed for private repos / higher rate limits)
        GEMINI_API_KEY - Free API key from https://aistudio.google.com/apikey

Usage:
    python generate_release_docs.py \
        --owner myorg \
        --repo myrepo \
        --tag1 v1.2.0 \
        --tag2 v1.3.0 \
        --prev-doc ./previous_release_notes.md \
        --output ./release_notes_v1.3.0.md

    This writes both release_notes_v1.3.0.md and release_notes_v1.3.0.docx.
"""

import argparse
import os
import re
import sys
import textwrap
import requests
from docx import Document
from docx.shared import Pt

GITHUB_API = "https://api.github.com"
GEMINI_API = "https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"

MAX_PATCH_CHARS_PER_FILE = 4000   # truncate huge file diffs
MAX_TOTAL_DIFF_CHARS = 120_000    # keep total prompt reasonable

BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*")


def _add_runs_with_bold(paragraph, text):
    """Split on **bold** markers and add runs preserving bold formatting."""
    pos = 0
    for m in BOLD_PATTERN.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        run = paragraph.add_run(m.group(1))
        run.bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def markdown_to_docx(markdown_text, output_path, title=None):
    """
    Minimal markdown -> docx converter covering what release notes typically use:
    headings (#, ##, ###), bullet lists (-, *), bold (**text**), and plain paragraphs.
    Good enough for AI-generated release notes; not a full CommonMark implementation.
    """
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Group source lines into blocks: blank lines separate blocks; soft-wrapped
    # lines within a plain paragraph get joined with a space.
    raw_lines = markdown_text.splitlines()
    blocks = []
    current = []
    for raw_line in raw_lines:
        line = raw_line.rstrip()
        if not line.strip():
            if current:
                blocks.append(current)
                current = []
            continue
        is_special = (
            line.strip() in ("---", "***", "___")
            or re.match(r"^(#{1,4})\s+", line)
            or re.match(r"^\s*[-*]\s+", line)
            or re.match(r"^\s*\d+\.\s+", line)
        )
        if is_special:
            if current:
                blocks.append(current)
                current = []
            blocks.append([line])
        else:
            current.append(line)
    if current:
        blocks.append(current)

    for block in blocks:
        line = block[0] if len(block) == 1 else " ".join(l.strip() for l in block)

        if line.strip() in ("---", "***", "___"):
            doc.add_paragraph("―" * 20)
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            doc.add_heading(text, level=level)
            continue

        bullet_match = re.match(r"^\s*[-*]\s+(.*)", line)
        if bullet_match:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, bullet_match.group(1).strip())
            continue

        numbered_match = re.match(r"^\s*\d+\.\s+(.*)", line)
        if numbered_match:
            p = doc.add_paragraph(style="List Number")
            _add_runs_with_bold(p, numbered_match.group(1).strip())
            continue

        # plain paragraph (possibly joined from soft-wrapped lines)
        p = doc.add_paragraph()
        _add_runs_with_bold(p, line.strip())

    doc.save(output_path)


def fetch_compare(owner, repo, tag1, tag2, token=None):
    url = f"{GITHUB_API}/repos/{owner}/{repo}/compare/{tag1}...{tag2}"
    headers = {"Accept": "application/vnd.github+json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    resp = requests.get(url, headers=headers, timeout=30)
    if resp.status_code != 200:
        sys.exit(f"GitHub API error ({resp.status_code}): {resp.text[:500]}")
    return resp.json()


def build_diff_summary(compare_data):
    lines = []
    lines.append(f"Total commits: {compare_data.get('total_commits', 'unknown')}")
    lines.append("")
    lines.append("## Commits")
    for c in compare_data.get("commits", []):
        msg = c["commit"]["message"].split("\n")[0]
        author = c["commit"]["author"]["name"]
        sha = c["sha"][:7]
        lines.append(f"- [{sha}] {msg} (by {author})")

    lines.append("")
    lines.append("## Changed files")
    total_chars = sum(len(l) for l in lines)
    for f in compare_data.get("files", []):
        header = f"\n### {f['filename']} ({f['status']}, +{f.get('additions', 0)}/-{f.get('deletions', 0)})\n"
        patch = f.get("patch", "")[:MAX_PATCH_CHARS_PER_FILE]
        block = header + "```diff\n" + patch + "\n```\n"
        if total_chars + len(block) > MAX_TOTAL_DIFF_CHARS:
            lines.append("\n...[remaining file diffs truncated for length]...")
            break
        lines.append(block)
        total_chars += len(block)

    return "\n".join(lines)


def generate_release_doc(diff_summary, prev_doc_text, tag1, tag2, api_key, model="gemini-2.0-flash"):
    prompt = textwrap.dedent(f"""
        You are a release notes writer. You will be given:
        1. A PREVIOUS RELEASE DOC that shows the exact structure, tone, section headers,
           and formatting style this project uses for release notes.
        2. A DIFF SUMMARY (commit messages + changed files) between two tags: {tag1} -> {tag2}.

        Your task: write a NEW release doc for the {tag1} -> {tag2} change, matching the
        structure/style/tone of the previous doc as closely as possible. Group changes into
        sensible sections (e.g. Features, Fixes, Breaking Changes, Chores) based on the commit
        messages and diffs — infer categories the same way the previous doc appears to.
        Do not invent changes that aren't supported by the commits/diff. If something is unclear,
        summarize conservatively rather than guessing details.

        --- PREVIOUS RELEASE DOC (style reference) ---
        {prev_doc_text}

        --- DIFF SUMMARY ({tag1} -> {tag2}) ---
        {diff_summary}

        Now write the new release doc in markdown, ready to publish.
    """).strip()

    url = GEMINI_API.format(model=model)
    resp = requests.post(
        url,
        params={"key": api_key},
        json={"contents": [{"parts": [{"text": prompt}]}]},
        timeout=120,
    )
    if resp.status_code != 200:
        sys.exit(f"Gemini API error ({resp.status_code}): {resp.text[:800]}")

    data = resp.json()
    try:
        return data["candidates"][0]["content"]["parts"][0]["text"]
    except (KeyError, IndexError):
        sys.exit(f"Unexpected Gemini response shape: {data}")


def main():
    parser = argparse.ArgumentParser(description="Generate release docs from a GitHub tag diff using Gemini.")
    parser.add_argument("--owner", required=True, help="GitHub repo owner/org")
    parser.add_argument("--repo", required=True, help="GitHub repo name")
    parser.add_argument("--tag1", required=True, help="Previous tag")
    parser.add_argument("--tag2", required=True, help="New tag")
    parser.add_argument("--prev-doc", required=True, help="Path to previous release doc (markdown/txt)")
    parser.add_argument("--output", default="release_notes.md", help="Output file path")
    parser.add_argument("--github-token", default=os.environ.get("GITHUB_TOKEN"))
    parser.add_argument("--gemini-key", default=os.environ.get("GEMINI_API_KEY"))
    parser.add_argument("--model", default="gemini-2.0-flash")
    args = parser.parse_args()

    if not args.gemini_key:
        sys.exit("Missing Gemini API key. Set GEMINI_API_KEY env var or pass --gemini-key.\n"
                  "Get a free key at https://aistudio.google.com/apikey")

    if not os.path.exists(args.prev_doc):
        sys.exit(f"Previous release doc not found: {args.prev_doc}")

    print(f"Fetching diff between {args.tag1} and {args.tag2} ...")
    compare_data = fetch_compare(args.owner, args.repo, args.tag1, args.tag2, args.github_token)

    print("Building diff summary ...")
    diff_summary = build_diff_summary(compare_data)

    with open(args.prev_doc, "r", encoding="utf-8") as f:
        prev_doc_text = f.read()

    print("Generating release doc with Gemini ...")
    output_text = generate_release_doc(
        diff_summary, prev_doc_text, args.tag1, args.tag2, args.gemini_key, args.model
    )

    with open(args.output, "w", encoding="utf-8") as f:
        f.write(output_text)

    docx_path = os.path.splitext(args.output)[0] + ".docx"
    markdown_to_docx(output_text, docx_path)

    print(f"Done. Release docs written to:\n  {args.output}\n  {docx_path}")


if __name__ == "__main__":
    main()
