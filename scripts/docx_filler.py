"""
docx_filler.py

Fills a .docx template that contains placeholder tokens:
  Scalar placeholders (plain text swap): {{VERSION}}, {{DATE}}, {{OVERVIEW}},
      {{CONTRIBUTORS}}, {{DIFF_LINK}}
  Table row placeholder (duplicated once per change): a row whose cells contain
      {{TYPE}}, {{DESCRIPTION}}, {{MODULE}}, {{OWNER}}
  List item placeholders (duplicated once per list item): a paragraph containing
      {{BREAKING_CHANGE_ITEM}} or {{KNOWN_ISSUE_ITEM}}

This preserves the template's fonts, colors, table borders, and paragraph styles —
only the text content changes.
"""

import copy
from docx import Document
from docx.oxml.ns import qn


def _set_paragraph_text(paragraph, text):
    """Replace a paragraph's visible text while keeping the first run's formatting."""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    first_run = paragraph.runs[0]
    first_run.text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def _replace_scalar_placeholders(doc, mapping):
    """Simple {{KEY}} -> value swap across all top-level body paragraphs."""
    for paragraph in doc.paragraphs:
        full_text = "".join(run.text for run in paragraph.runs)
        if not full_text:
            continue
        replaced = full_text
        changed = False
        for key, value in mapping.items():
            token = "{{" + key + "}}"
            if token in replaced:
                replaced = replaced.replace(token, value)
                changed = True
        if changed:
            _set_paragraph_text(paragraph, replaced)


def _find_table_with_placeholder(doc, marker_token):
    for table in doc.tables:
        for row in table.rows:
            for c in row.cells:
                if marker_token in c.text:
                    return table, row
    return None, None


def _expand_table_rows(doc, changes):
    """
    Find the row containing {{TYPE}}/{{DESCRIPTION}}/{{MODULE}}/{{OWNER}} and
    duplicate it once per entry in `changes`, then remove the template row.
    """
    table, template_row = _find_table_with_placeholder(doc, "{{TYPE}}")
    if template_row is None:
        return  # template has no table row placeholder; nothing to do

    template_tr = template_row._tr
    anchor = template_tr

    for change in changes:
        new_tr = copy.deepcopy(template_tr)
        anchor.addnext(new_tr)
        anchor = new_tr

        from docx.table import _Row
        new_row = _Row(new_tr, table)
        mapping = {
            "{{TYPE}}": change.get("type", ""),
            "{{DESCRIPTION}}": change.get("description", ""),
            "{{MODULE}}": change.get("module", ""),
            "{{OWNER}}": change.get("owner", ""),
        }
        for cell_obj in new_row.cells:
            for p in cell_obj.paragraphs:
                text = "".join(r.text for r in p.runs)
                for token, val in mapping.items():
                    if token in text:
                        text = text.replace(token, val)
                if text != "".join(r.text for r in p.runs):
                    _set_paragraph_text(p, text)

    # remove the original placeholder row
    template_tr.getparent().remove(template_tr)


def _find_paragraph_with_token(doc, token):
    for p in doc.paragraphs:
        if token in p.text:
            return p
    return None


def _expand_list_placeholder(doc, token, items, fallback_text="None."):
    """
    Find the paragraph containing `token`, duplicate it once per item in `items`
    (preserving bullet/paragraph formatting), fill each with the item text, then
    remove the original placeholder paragraph. If `items` is empty, the paragraph
    is replaced with `fallback_text` instead of being removed.
    """
    template_p = _find_paragraph_with_token(doc, token)
    if template_p is None:
        return

    template_elem = template_p._p

    if not items:
        _set_paragraph_text(template_p, fallback_text)
        return

    anchor = template_elem
    for item in items:
        new_elem = copy.deepcopy(template_elem)
        anchor.addnext(new_elem)
        anchor = new_elem

        from docx.text.paragraph import Paragraph
        new_p = Paragraph(new_elem, template_p._parent)
        text = "".join(r.text for r in new_p.runs)
        text = text.replace(token, item)
        _set_paragraph_text(new_p, text)

    # remove the original placeholder paragraph
    template_elem.getparent().remove(template_elem)


def fill_release_template(template_path, output_path, data):
    """
    data: {
        "version": str, "date": str, "overview": str,
        "changes": [{"type","description","module","owner"}, ...],
        "breaking_changes": [str, ...],
        "known_issues": [str, ...],
        "contributors": [str, ...],  # will be joined with ", "
        "diff_link": str,
    }
    """
    doc = Document(template_path)

    _replace_scalar_placeholders(doc, {
        "VERSION": data.get("version", ""),
        "DATE": data.get("date", ""),
        "OVERVIEW": data.get("overview", ""),
        "CONTRIBUTORS": ", ".join(data.get("contributors", [])) or "N/A",
        "DIFF_LINK": data.get("diff_link", ""),
    })

    _expand_table_rows(doc, data.get("changes", []))
    _expand_list_placeholder(doc, "{{BREAKING_CHANGE_ITEM}}", data.get("breaking_changes", []), fallback_text="None in this release.")
    _expand_list_placeholder(doc, "{{KNOWN_ISSUE_ITEM}}", data.get("known_issues", []), fallback_text="None reported.")

    doc.save(output_path)
